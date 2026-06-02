import numpy as np
from PIL import Image
import os
import torch
import warnings
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
import cv2
import kornia

from scipy.signal import convolve2d
import math
import torch.nn.functional as F

def QNCIE_function(ir_img_tensor, vi_img_tensor, f_img_tensor):
    def normalize1(img_tensor):
        img_min = img_tensor.min()
        img_max = img_tensor.max()
        return (img_tensor - img_min) / (img_max - img_min)
    def NCC(img1, img2):
        mean1 = torch.mean(img1)
        mean2 = torch.mean(img2)
        numerator = torch.sum((img1 - mean1) * (img2 - mean2))
        denominator = torch.sqrt(torch.sum((img1 - mean1) ** 2) * torch.sum((img2 - mean2) ** 2))
        return numerator / (denominator + 1e-10)

    ir_img_tensor = normalize1(ir_img_tensor)
    vi_img_tensor = normalize1(vi_img_tensor)
    f_img_tensor = normalize1(f_img_tensor)

    NCCxy = NCC(ir_img_tensor, vi_img_tensor)
    NCCxf = NCC(ir_img_tensor, f_img_tensor)
    NCCyf = NCC(vi_img_tensor, f_img_tensor)
    R = torch.tensor([[1, NCCxy, NCCxf],
                      [NCCxy, 1, NCCyf],
                      [NCCxf, NCCyf, 1]], dtype=torch.float32)

    r = torch.linalg.eigvals(R).real
    K = 3
    b = 256
    HR = torch.sum(r * torch.log2(r / K) / K)
    HR = -HR / np.log2(b)
    QNCIE = 1 - HR.item()
    return QNCIE

def fspecial_gaussian(shape, sigma):
    m, n = [(ss-1.)/2. for ss in shape]
    y, x = np.ogrid[-m:m+1, -n:n+1]
    h = np.exp(-(x*x + y*y) / (2.*sigma*sigma))
    h[h < np.finfo(h.dtype).eps*h.max()] = 0
    sumh = h.sum()
    if sumh != 0:
        h /= sumh
    return h

def fspecial_gaussian(size, sigma):
    x = torch.linspace(-size[0]//2, size[0]//2, size[0])
    y = torch.linspace(-size[1]//2, size[1]//2, size[1])
    x, y = torch.meshgrid(x, y)
    g = torch.exp(-(x**2 + y**2) / (2 * sigma**2))
    return g / g.sum()

def convolve2d_torch(input, kernel):
    kernel = kernel.unsqueeze(0).unsqueeze(0).to(input.device)  # Add batch and channel dimensions
    return F.conv2d(input.unsqueeze(0).unsqueeze(0), kernel, padding=kernel.shape[2] // 2)[0][0]

def vifp_mscale(ref, dist):
    sigma_nsq = 2
    num = 0
    den = 0
    for scale in range(1, 5):
        N = 2 ** (4 - scale + 1) + 1
        win = fspecial_gaussian((N, N), N / 5)

        if scale > 1:
            ref = convolve2d_torch(ref, win)
            dist = convolve2d_torch(dist, win)
            ref = ref[::2, ::2]
            dist = dist[::2, ::2]

        mu1 = convolve2d_torch(ref, win)
        mu2 = convolve2d_torch(dist, win)
        mu1_sq = mu1 * mu1
        mu2_sq = mu2 * mu2
        mu1_mu2 = mu1 * mu2
        sigma1_sq = convolve2d_torch(ref * ref, win) - mu1_sq
        sigma2_sq = convolve2d_torch(dist * dist, win) - mu2_sq
        sigma12 = convolve2d_torch(ref * dist, win) - mu1_mu2
        sigma1_sq[sigma1_sq < 0] = 0
        sigma2_sq[sigma2_sq < 0] = 0

        g = sigma12 / (sigma1_sq + 1e-10)
        sv_sq = sigma2_sq - g * sigma12

        g[sigma1_sq < 1e-10] = 0
        sv_sq[sigma1_sq < 1e-10] = sigma2_sq[sigma1_sq < 1e-10]
        sigma1_sq[sigma1_sq < 1e-10] = 0

        g[sigma2_sq < 1e-10] = 0
        sv_sq[sigma2_sq < 1e-10] = 0

        sv_sq[g < 0] = sigma2_sq[g < 0]
        g[g < 0] = 0
        sv_sq[sv_sq <= 1e-10] = 1e-10

        num += torch.sum(torch.log10(1 + g**2 * sigma1_sq / (sv_sq + sigma_nsq)))
        den += torch.sum(torch.log10(1 + sigma1_sq / sigma_nsq))

    vifp = num / den
    return vifp

def VIF_function(A, B, F):
    VIF = vifp_mscale(A, F) + vifp_mscale(B, F)
    return VIF

def sobel_fn(x):
    vtemp = np.array([[-1, 0, 1], [-2, 0, 2], [-1, 0, 1]]) / 8
    htemp = np.array([[-1, -2, -1], [0, 0, 0], [1, 2, 1]]) / 8

    a, b = htemp.shape
    x_ext = per_extn_im_fn(x, a)
    p, q = x_ext.shape
    gv = np.zeros((p - 2, q - 2))
    gh = np.zeros((p - 2, q - 2))
    gv = convolve2d(x_ext, vtemp, mode='valid')
    gh = convolve2d(x_ext, htemp, mode='valid')

    return gv, gh


def per_extn_im_fn(x, wsize):
    hwsize = (wsize - 1) // 2

    p, q = x.shape
    xout_ext = np.zeros((p + wsize - 1, q + wsize - 1))
    xout_ext[hwsize: p + hwsize, hwsize: q + hwsize] = x


    if wsize - 1 == hwsize + 1:
        xout_ext[0: hwsize, :] = xout_ext[2, :].reshape(1, -1)
        xout_ext[p + hwsize: p + wsize - 1, :] = xout_ext[-3, :].reshape(1, -1)

    xout_ext[:, 0: hwsize] = xout_ext[:, 2].reshape(-1, 1)
    xout_ext[:, q + hwsize: q + wsize - 1] = xout_ext[:, -3].reshape(-1, 1)

    return xout_ext

def get_Qabf(pA, pB, pF):
    L = 1
    Tg = 0.9994
    kg = -15
    Dg = 0.5;
    Ta = 0.9879
    ka = -22
    Da = 0.8

    h1 = np.array([[1, 2, 1], [0, 0, 0], [-1, -2, -1]]).astype(np.float32)
    h2 = np.array([[0, 1, 2], [-1, 0, 1], [-2, -1, 0]]).astype(np.float32)
    h3 = np.array([[-1, 0, 1], [-2, 0, 2], [-1, 0, 1]]).astype(np.float32)


    strA = pA
    strB = pB
    strF = pF

    def flip180(arr):
        return np.flip(arr)

    def convolution(k, data):
        k = flip180(k)
        data = np.pad(data, ((1, 1), (1, 1)), 'constant', constant_values=(0, 0))
        img_new = convolve2d(data, k, mode='valid')
        return img_new

    def getArray(img):
        SAx = convolution(h3, img)
        SAy = convolution(h1, img)
        gA = np.sqrt(np.multiply(SAx, SAx) + np.multiply(SAy, SAy))
        n, m = img.shape
        aA = np.zeros((n, m))
        zero_mask = SAx == 0
        aA[~zero_mask] = np.arctan(SAy[~zero_mask] / SAx[~zero_mask])
        aA[zero_mask] = np.pi / 2
        return gA, aA

    gA, aA = getArray(strA)
    gB, aB = getArray(strB)
    gF, aF = getArray(strF)

    def getQabf(aA, gA, aF, gF):
        mask = (gA > gF)
        GAF = np.where(mask, gF / gA, np.where(gA == gF, gF, gA / gF))

        AAF = 1 - np.abs(aA - aF) / (math.pi / 2)

        QgAF = Tg / (1 + np.exp(kg * (GAF - Dg)))
        QaAF = Ta / (1 + np.exp(ka * (AAF - Da)))

        QAF = QgAF * QaAF
        return QAF

    QAF = getQabf(aA, gA, aF, gF)
    QBF = getQabf(aB, gB, aF, gF)

    deno = np.sum(gA + gB)
    nume = np.sum(np.multiply(QAF, gA) + np.multiply(QBF, gB))
    output = nume / deno
    return output

def Qabf_function(A, B, F):
    return get_Qabf(A, B, F)

def Hab(im1, im2, gray_level):
	hang, lie = im1.shape
	count = hang * lie
	N = gray_level
	h = np.zeros((N, N))
	for i in range(hang):
		for j in range(lie):
			h[im1[i, j], im2[i, j]] = h[im1[i, j], im2[i, j]] + 1
	h = h / np.sum(h)
	im1_marg = np.sum(h, axis=0)
	im2_marg = np.sum(h, axis=1)
	H_x = 0
	H_y = 0
	for i in range(N):
		if (im1_marg[i] != 0):
			H_x = H_x + im1_marg[i] * math.log2(im1_marg[i])
	for i in range(N):
		if (im2_marg[i] != 0):
			H_x = H_x + im2_marg[i] * math.log2(im2_marg[i])
	H_xy = 0
	for i in range(N):
		for j in range(N):
			if (h[i, j] != 0):
				H_xy = H_xy + h[i, j] * math.log2(h[i, j])
	MI = H_xy - H_x - H_y
	return MI

def MI_function(A, B, F, gray_level=256):
	MIA = Hab(A, F, gray_level)
	MIB = Hab(B, F, gray_level)
	MI_results = MIA + MIB
	return MI_results

warnings.filterwarnings("ignore")
device = torch.device("cuda" if torch.cuda.is_available() else "cpu")


def write_excel(excel_name='metric.xlsx', worksheet_name='VIF', column_index=0, data=None):
    try:
        workbook = load_workbook(excel_name)
    except FileNotFoundError:
        workbook = Workbook()

    worksheet = workbook.create_sheet(title=worksheet_name) if worksheet_name not in workbook.sheetnames else workbook[
        worksheet_name]

    column = get_column_letter(column_index + 1)
    for i, value in enumerate(data):
        cell = worksheet[column + str(i + 1)]
        cell.value = value

    workbook.save(excel_name)


def evaluation_one(ir_name, vi_name, f_name):
    f_img = Image.open(f_name).convert('L')
    ir_img = Image.open(ir_name).convert('L')
    vi_img = Image.open(vi_name).convert('L')

    vi_img = vi_img.resize((f_img.size[0], f_img.size[1]), Image.BICUBIC)

    f_img_tensor = torch.tensor(np.array(f_img)).float().to(device)
    ir_img_tensor = torch.tensor(np.array(ir_img)).float().to(device)
    vi_img_tensor = torch.tensor(np.array(vi_img)).float().to(device)

    f_img_int = np.array(f_img).astype(np.int32)
    f_img_double = np.array(f_img).astype(np.float32)

    ir_img_int = np.array(ir_img).astype(np.int32)
    ir_img_double = np.array(ir_img).astype(np.float32)

    vi_img_int = np.array(vi_img).astype(np.int32)
    vi_img_double = np.array(vi_img).astype(np.float32)


    QNCIE = QNCIE_function(ir_img_tensor, vi_img_tensor, f_img_tensor)
    MI = MI_function(ir_img_int, vi_img_int, f_img_int, gray_level=256)
    VIF = VIF_function(ir_img_tensor, vi_img_tensor, f_img_tensor)
    Qabf = Qabf_function(ir_img_double, vi_img_double, f_img_double)

    return QNCIE, MI, VIF, Qabf

def imread(path, flags=cv2.IMREAD_GRAYSCALE, unsqueeze=False):
    im_cv = cv2.imread(str(path), flags)
    assert im_cv is not None, f"Image {str(path)} is invalid."
    im_ts = kornia.utils.image_to_tensor(im_cv / 255.).type(torch.FloatTensor)
    return im_ts.unsqueeze(0) if unsqueeze else im_ts


if __name__ == '__main__':
    if __name__ == '__main__':
        with_mean = True

        
        QNCIE_list = []
        MI_list = []
        VIF_list = []
        Qabf_list = []

        root_ir = './Results/MMMF_adjust/TNO/Final_1200/reg'
        root_vis = './dataset/TNO/vi'
        root_fus = './Results/MMMF_adjust/TNO/Final_1200/fus'

        vi_img_list = sorted(os.listdir(root_ir))
        ir_img_list = sorted(os.listdir(root_vis))
        fuse_img_list = sorted(os.listdir(root_fus))


    for vi, ir,fuse in zip(vi_img_list, ir_img_list,fuse_img_list):
        vi_img_path = os.path.join(root_vis,vi)
        ir_img_path = os.path.join(root_ir, ir)
        fus_img_path = os.path.join(root_fus, fuse)

        QNCIE, MI, VIF, Qabf = evaluation_one(ir_img_path, vi_img_path, fus_img_path)
        QNCIE_list.append(QNCIE)
        MI_list.append(MI)
        VIF_list.append(VIF)
        Qabf_list.append(Qabf)

    if with_mean:
        QNCIE_tensor = torch.tensor(QNCIE_list).mean().item()
        print('QNCIE:{}'.format(QNCIE_tensor))
        MI_tensor = torch.tensor(MI_list).mean().item()
        print('MI:{}'.format(MI_tensor))
        VIF_tensor = torch.tensor(VIF_list).mean().item()
        print('VIF:{}'.format(VIF_tensor))
        Qabf_tensor = torch.tensor(Qabf_list).mean().item()
        print('Qabf:{}'.format(Qabf_tensor))